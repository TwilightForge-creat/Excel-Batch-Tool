# -*- coding: utf-8 -*-
"""
批量处理 Excel 文件 —— 通用模板
====================================
功能简介：
1. 递归遍历指定输入文件夹下的所有 .xlsx 文件；
2. 使用 openpyxl 获取每个工作簿中的所有工作表名称；
3. 使用 pandas 读取每个工作表，并进行基础清洗（删除完全空行、重置索引等）；
4. 将处理后的内容保存到新的输出文件夹，绝不覆盖原文件；
5. 每个文件都使用 try...except 单独捕获异常，单个文件损坏不会导致整个程序崩溃；
6. 在终端打印详细处理日志。
作者：熔壑工坊
"""

import logging
# =====================================================================
# 【可修改区域 2】输出文件夹路径
# 说明：处理后的文件会保存到这个文件夹中，原文件不会被修改或覆盖。
#       如果该文件夹不存在，程序会自动创建。
# =====================================================================
# OUTPUT_FOLDER = r"输出文件夹路径"
# 傻瓜版
import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# =====================================================================
# 【可修改区域 1】输入文件夹路径
# 说明：脚本会递归遍历该文件夹下的所有 .xlsx 文件（包括子文件夹）。
# 注意：
#   1. 路径前加 r 是为了防止 Windows 路径中的反斜杠被转义；
#   2. 也可以使用正斜杠。
# =====================================================================
# INPUT_FOLDER = r"输入文件夹路径"

# 核心逻辑：判断当前是在 Python 环境跑，还是打包成了 EXE
if getattr(sys, 'frozen', False):
    # 如果是 EXE，获取 EXE 文件所在的真实目录
    BASE_DIR = os.path.dirname(sys.executable)
else:
    # 如果是 Python 脚本，获取脚本所在的目录
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# 拼接出 input 和 output 的绝对路径
INPUT_FOLDER = os.path.join(BASE_DIR, "input")
OUTPUT_FOLDER = os.path.join(BASE_DIR, "output")

# =====================================================================
# 【可修改区域 3】文件名关键字过滤（可选）
# 说明：默认处理所有 .xlsx 文件。
#       如果你只想处理文件名中包含某个关键字的文件，请修改 KEYWORD。
#       例如：KEYWORD = "工资"  表示只处理文件名中包含“工资”的 Excel 文件。
#       如果不需要过滤，请保持 KEYWORD = None。
# =====================================================================
KEYWORD = None


# =====================================================================
# 【可修改区域 4】是否保留原始文件夹结构
# 说明：
#   True  -> 输出文件夹中会按照输入文件夹的相对目录结构保存文件；
#   False -> 所有处理后的文件都直接存放在输出文件夹根目录下。
# =====================================================================
KEEP_FOLDER_STRUCTURE = True


# 配置日志格式：终端打印时间、级别和具体信息
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)


def get_sheet_names(file_path: Path) -> list:
    """
    使用 openpyxl 获取工作簿中的所有工作表名称。

    参数:
        file_path: Excel 文件路径

    返回:
        工作表名称列表，例如 ["Sheet1", "Sheet2"]
    """
    # read_only=True：只读模式，内存占用更低，速度更快
    # data_only=True：只读取单元格结果，不读取公式本身
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        # 无论是否发生异常，都要关闭工作簿，释放资源
        workbook.close()


def is_under_output_folder(file_path: Path, output_dir: Path) -> bool:
    """
    判断某个文件是否位于输出文件夹内部。
    用于避免输出文件夹被放在输入文件夹内部时，程序重复处理已经生成的结果文件。

    参数:
        file_path: 待判断的文件路径
        output_dir: 输出文件夹路径

    返回:
        如果 file_path 位于 output_dir 内部，则返回 True，否则返回 False。
    """
    try:
        file_path.resolve().relative_to(output_dir.resolve())
        return True
    except ValueError:
        return False


def process_excel_file(src_path: Path, dst_path: Path) -> None:
    """
    处理单个 Excel 文件的通用模板。

    处理流程：
        1. 使用 openpyxl 读取工作簿中的所有工作表名称；
        2. 使用 pandas 逐表读取数据；
        3. 对每个工作表执行基础清洗；
        4. 将所有处理后的工作表写入新的 Excel 文件。

    参数:
        src_path: 原始 Excel 文件路径
        dst_path: 处理后 Excel 文件输出路径
    """
    logging.info("  正在获取工作表信息：%s", src_path.name)

    # 第一步：获取所有工作表名称
    sheet_names = get_sheet_names(src_path)
    if not sheet_names:
        raise ValueError("工作簿中不包含任何工作表")

    # 第二步：使用 pandas 读取整个工作簿
    # 指定 engine="openpyxl"，确保 pandas 使用 openpyxl 处理 .xlsx 文件
    xls = pd.ExcelFile(src_path, engine="openpyxl")

    # 用字典保存处理后的数据：{工作表名称: DataFrame}
    processed_sheets = {}

    try:
        for sheet_name in sheet_names:
            logging.info("  正在处理工作表：'%s'", sheet_name)

            # 读取当前工作表
            df = xls.parse(sheet_name=sheet_name)

            # ---------------------------------------------------------
            # 基础处理 1：删除完全空白的行
            # dropna(how="all") 表示：当一行中的所有单元格都为空时，删除该行。
            # ---------------------------------------------------------
            df = df.dropna(how="all")

            # ---------------------------------------------------------
            # 基础处理 2：重置索引
            # 删除空行后，索引可能变得不连续，例如 0, 2, 5……
            # reset_index(drop=True) 会把索引重新排列为 0, 1, 2……
            # drop=True 表示不把旧索引保存为新列。
            # ---------------------------------------------------------
            df = df.reset_index(drop=True)

            # ---------------------------------------------------------
            # 基础处理 3（可选）：去除字符串列首尾空格
            # 如果不需要，可以保留注释状态或直接删除。
            # ---------------------------------------------------------
            # for col in df.columns:
            #     if df[col].dtype == "object":
            #         df[col] = df[col].map(
            #             lambda x: x.strip() if isinstance(x, str) else x
            #         )

            # =========================================================
            # 【可修改区域 5】核心业务处理逻辑
            # 说明：在这里添加你自己的 Excel 处理逻辑。
            # 下面提供了几个常见示例，请按需取消注释并修改。
            # =========================================================

            # ---------------------------------------------------------
            # 示例 1：只保留需要的列
            # ---------------------------------------------------------
            # keep_columns = ["姓名", "部门", "基本工资"]
            # existing_columns = [c for c in keep_columns if c in df.columns]
            # if existing_columns:
            #     df = df[existing_columns]

            # ---------------------------------------------------------
            # 示例 2：删除关键列为空的行
            # ---------------------------------------------------------
            # key_column = "姓名"
            # if key_column in df.columns:
            #     df = df.dropna(subset=[key_column])

            # ---------------------------------------------------------
            # 示例 3：新增计算列
            # ---------------------------------------------------------
            # if "单价" in df.columns and "数量" in df.columns:
            #     df["总价"] = df["单价"] * df["数量"]

            # ---------------------------------------------------------
            # 示例 4：按某列排序
            # ---------------------------------------------------------
            # sort_column = "日期"
            # if sort_column in df.columns:
            #     df = df.sort_values(by=sort_column)

            # ---------------------------------------------------------
            # 示例 5：删除重复行
            # ---------------------------------------------------------
            # df = df.drop_duplicates()

            # ---------------------------------------------------------
            # 示例 6：筛选满足条件的行
            # ---------------------------------------------------------
            # if "工资" in df.columns:
            #     df = df[df["工资"] > 5000]

            # =========================================================
            # 【可修改区域 5】结束
            # =========================================================

            # 保存处理后的工作表
            processed_sheets[sheet_name] = df
    finally:
        # 无论处理是否成功，都要关闭 ExcelFile，释放资源
        xls.close()

    # 第三步：将处理后的所有工作表写入新的 Excel 文件
    # 使用 pd.ExcelWriter 并指定 engine="openpyxl"
    with pd.ExcelWriter(dst_path, engine="openpyxl") as writer:
        for sheet_name, df in processed_sheets.items():
            # index=False 表示不把 DataFrame 的行索引写入 Excel
            df.to_excel(writer, sheet_name=sheet_name, index=False)


def main() -> None:
    """
    主函数：批量处理所有符合条件的 Excel 文件。
    """
    input_dir = Path(INPUT_FOLDER)
    output_dir = Path(OUTPUT_FOLDER)

    # 检查输入文件夹是否存在
    if not input_dir.exists():
        logging.error("输入文件夹不存在：%s", INPUT_FOLDER)
        sys.exit(1)

    # 创建输出文件夹。
    # parents=True 表示如果多级目录不存在都一并创建；
    # exist_ok=True 表示如果目录已存在则不会报错。
    output_dir.mkdir(parents=True, exist_ok=True)

    # ---------------------------------------------------------
    # 查找所有需要处理的 .xlsx 文件
    # 1. rglob("*") 递归遍历所有文件；
    # 2. suffix.lower() == ".xlsx" 自动兼容 .xlsx 和 .XLSX；
    # 3. 跳过文件名以 ~$ 开头的 Office 临时锁定文件；
    # 4. 如果输出文件夹位于输入文件夹内部，则跳过输出文件夹中的文件。
    # ---------------------------------------------------------
    all_files = [
        p
        for p in input_dir.rglob("*")
        if p.is_file()
        and p.suffix.lower() == ".xlsx"
        and not p.name.startswith("~$")
        and not is_under_output_folder(p, output_dir)
    ]

    # 根据关键字过滤文件名
    if KEYWORD:
        all_files = [p for p in all_files if KEYWORD in p.stem]

    if not all_files:
        logging.warning("在输入文件夹中没有找到任何符合要求的 .xlsx 文件。")
        return

    logging.info("共找到 %d 个 Excel 文件，开始批量处理……", len(all_files))

    success_count = 0
    fail_count = 0

    for src_path in all_files:
        # ---------------------------------------------------------
        # 决定输出文件路径
        # 如果 KEEP_FOLDER_STRUCTURE = True，则保留相对目录结构；
        # 否则直接放在输出文件夹根目录。
        # ---------------------------------------------------------
        if KEEP_FOLDER_STRUCTURE:
            relative_path = src_path.relative_to(input_dir)
            dst_path = output_dir / relative_path
        else:
            dst_path = output_dir / src_path.name

        # 确保输出文件所在目录存在
        dst_path.parent.mkdir(parents=True, exist_ok=True)

        # ---------------------------------------------------------
        # try...except 异常处理：
        # 单个文件损坏或处理失败时，只记录错误并继续处理下一个文件，
        # 不会导致整个批量处理程序崩溃。
        # ---------------------------------------------------------
        try:
            logging.info("开始处理：%s", src_path)
            process_excel_file(src_path, dst_path)
            success_count += 1
            logging.info("处理完成：%s -> %s", src_path, dst_path)
        except Exception as exc:
            fail_count += 1
            # exc_info=True 会打印完整错误堆栈，便于排查问题
            logging.error("处理失败：%s，原因：%s", src_path, exc, exc_info=True)

    # 打印最终统计结果
    logging.info("=" * 60)
    logging.info("批量处理结束：成功 %d 个，失败 %d 个。", success_count, fail_count)
    logging.info("输出文件夹：%s", output_dir)


if __name__ == "__main__":
    main()